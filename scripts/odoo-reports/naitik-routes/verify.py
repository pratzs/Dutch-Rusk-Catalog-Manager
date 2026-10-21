"""Independent re-check of the Naitik route workbook. Fresh read from Odoo, own logic, no code shared with build.py.
usage: python -I verify.py WORKBOOK FY_START TODAY"""
import xmlrpc.client, ssl, os, sys, collections, datetime as dt
from zoneinfo import ZoneInfo
import openpyxl

WB, FY0, TODAY = sys.argv[1], dt.date.fromisoformat(sys.argv[2]), dt.date.fromisoformat(sys.argv[3])
env = {}
for l in open(os.path.expanduser("~/.odoo-worthy.env")):
    if "=" in l and not l.startswith("#"): k, v = l.strip().split("=", 1); env[k] = v.strip().strip('"')
ctx = ssl._create_unverified_context()
uid = xmlrpc.client.ServerProxy(env["ODOO_URL"] + "/xmlrpc/2/common", context=ctx).authenticate(env["ODOO_DB"], env["ODOO_USER"], env["ODOO_PASS"], {})
mo = xmlrpc.client.ServerProxy(env["ODOO_URL"] + "/xmlrpc/2/object", context=ctx, allow_none=True)
def x(m, meth, *a, **k): return mo.execute_kw(env["ODOO_DB"], uid, env["ODOO_PASS"], m, meth, list(a), k)
C = {"allowed_company_ids": [4]}
naitik = x("res.users", "search", [["name", "=", "Naitik Trivedi"]], context={"active_test": False})[0]
lines = x("account.move.line", "search_read", [["move_id.company_id", "=", 4], ["move_id.move_type", "in", ["out_invoice", "out_refund"]], ["move_id.state", "=", "posted"],
          ["move_id.invoice_date", ">=", str(FY0)], ["move_id.invoice_date", "<=", str(TODAY)], ["move_id.invoice_user_id", "=", naitik],
          ["display_type", "=", "product"], ["product_id", "!=", False]], fields=["move_id", "balance"], context=C)
val = collections.defaultdict(float)
for l in lines: val[l["move_id"][0]] -= l["balance"]
moves = x("account.move", "read", sorted({l["move_id"][0] for l in lines}), fields=["name", "move_type", "invoice_date", "invoice_origin", "partner_shipping_id"], context=C)
orig = sorted({m["invoice_origin"] for m in moves if m["invoice_origin"]})
so = {s["name"]: s for s in x("sale.order", "search_read", [["name", "in", orig], ["company_id", "=", 4]], fields=["name", "date_order"], context=C)}
NZ = ZoneInfo("Pacific/Auckland")
def odate(m):
    s = so.get(m["invoice_origin"] or "")
    return dt.datetime.strptime(s["date_order"], "%Y-%m-%d %H:%M:%S").replace(tzinfo=dt.timezone.utc).astimezone(NZ).date() if s else None
route = {}; usual = collections.defaultdict(lambda: [0.0, 0.0])
for m in moves:
    d = odate(m)
    if d and d.weekday() < 5:
        route[m["id"]] = "MW" if d.weekday() < 3 else "TF"
        if m["move_type"] == "out_invoice": usual[m["partner_shipping_id"][0]][0 if route[m["id"]] == "MW" else 1] += val[m["id"]]
for m in moves:
    if m["id"] not in route:
        u = usual.get(m["partner_shipping_id"][0])
        route[m["id"]] = ("MW" if u[0] >= u[1] else "TF") if u else "NONE"
def week(m):
    d = dt.date.fromisoformat(m["invoice_date"]) if m["move_type"] == "out_refund" else (odate(m) or dt.date.fromisoformat(m["invoice_date"]))
    return d - dt.timedelta(days=d.weekday())
first = FY0 + dt.timedelta(days=(7 - FY0.weekday()) % 7); last = TODAY - dt.timedelta(days=TODAY.weekday() + 7)
weeks = [first + dt.timedelta(weeks=i) for i in range((last - first).days // 7 + 1)]
wk = collections.defaultdict(float); tot = collections.defaultdict(float)
for m in moves:
    tot[route[m["id"]]] += val[m["id"]]; wk[(route[m["id"]], week(m))] += val[m["id"]]
avg = {r: sum(wk[(r, w)] for w in weeks) / len(weeks) for r in ("MW", "TF")}

wb = openpyxl.load_workbook(WB); fails = []
s = {r[0]: r[1:4] for r in wb["Summary"].iter_rows(min_row=5, max_row=12, values_only=True) if r[0]}
def near(a, b): return abs(a - b) < 0.01
if not (near(s["Average sales per week"][0], avg["MW"]) and near(s["Average sales per week"][1], avg["TF"])): fails.append("averages")
if not (near(s["Total sales this financial year"][0], tot["MW"]) and near(s["Total sales this financial year"][1], tot["TF"])): fails.append("totals")
rows = [r for r in wb["Weekly"].iter_rows(min_row=2, values_only=True) if r[0] and str(r[0]).startswith("Week ") and not str(r[0]).startswith("Week of")]
if len(rows) != len(weeks): fails.append("week count")
for i, r in enumerate(rows):
    w = weeks[i]
    if r[0] != f"Week {i + 1}" or r[1] != f"{w.day} {w:%b %Y}" or not near(r[2], wk[("MW", w)]) or not near(r[3], wk[("TF", w)]): fails.append(f"week {i + 1}")
for col, r_ in (("Mon-Wed", "MW"), ("Thu-Fri", "TF")):
    inv = [r for r in wb["Invoices " + col].iter_rows(min_row=2, values_only=True) if r[2]]
    names = {m["name"] for m in moves if route[m["id"]] == r_}
    if {r[2] for r in inv} != names: fails.append(col + " invoice list")
    if not near(sum(r[5] for r in inv), tot[r_]): fails.append(col + " invoice total")
    cu = [r for r in wb["Customers " + col].iter_rows(min_row=2, values_only=True) if r[0] and r[0] != "Total"]
    if not near(sum(r[3] for r in cu), tot[r_]): fails.append(col + " customer total")
grand = sum(val.values())
print(f"Odoo now: Naitik total ${grand:,.2f} | Mon-Wed ${tot['MW']:,.2f} avg ${avg['MW']:,.2f} | Thu-Fri ${tot['TF']:,.2f} avg ${avg['TF']:,.2f} | not on a route ${tot['NONE']:,.2f} | weeks {len(weeks)}")
print("FAILURES:", fails or 0)
if os.environ.get("SHOWDIFF"):
    for col, r_ in (("Mon-Wed", "MW"), ("Thu-Fri", "TF")):
        inv = {r[2]: r for r in wb["Invoices " + col].iter_rows(min_row=2, values_only=True) if r[2]}
        names = {m["name"] for m in moves if route[m["id"]] == r_}
        print(col, "only in workbook:", [(n, inv[n][5]) for n in sorted(set(inv) - names)][:10], "| only in Odoo check:", sorted(names - set(inv))[:10])
