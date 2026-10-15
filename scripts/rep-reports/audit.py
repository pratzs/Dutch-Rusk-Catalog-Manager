"""Independent audit of the finished rep workbooks (does not import the builder).
usage: python -I audit.py DATA_DIR OUT_DIR FROM TO
Needs OSTENDO_BASE_URL / OSTENDO_API_KEY in the environment for the live spot check."""
import json, glob, os, re, sys, random, datetime as dt, collections, urllib.request, ssl
import openpyxl

DATA, OUT = sys.argv[1], sys.argv[2]
FROM, TO = dt.date.fromisoformat(sys.argv[3]), dt.date.fromisoformat(sys.argv[4])
pd = lambda s: dt.datetime.strptime(s, "%d/%m/%Y").date()
H = json.load(open(DATA + "/hdr.json", encoding="utf-8"))
RAW = {h["INVOICENUMBER"]: h for h in H}
EV = {}
for l in open(DATA + "/oldsite.txt", encoding="utf-8").read().split(chr(10)):
    if not l: continue
    i, u, c, st, at, t = l.split("|")
    m = re.match(r"(\w+)\.? (\d+), (\d+)", at)
    EV["Old website order " + i] = (dt.datetime.strptime(f"{m.group(1)[:3]} {m.group(2)} {m.group(3)}", "%b %d %Y").date(), float(t) if t else None)
for o in json.load(open(DATA + "/shop.json", encoding="utf-8")):
    d = dt.datetime.fromisoformat(o["createdAt"].replace("Z", "+00:00")).astimezone(dt.timezone(dt.timedelta(hours=13))).date()
    EV["Shopify " + o["name"]] = (d, float(o["currentSubtotalPriceSet"]["shopMoney"]["amount"]))

fails = []
def check(cond, msg):
    if not cond: fails.append(msg)

raw = collections.defaultdict(lambda: [0, 0.0, 0, 0.0, set()])
for h in H:
    if FROM <= pd(h["INVOICEDATE"]) <= TO:
        r = raw[h["SALESPERSON"].split("-")[0]]; r[4].add(h["CUSTOMER"])
        if h["INVOICEORCREDIT"] == "Invoice": r[0] += 1; r[1] += h["INVOICENETTAMOUNT"]
        else: r[2] += 1; r[3] += h["INVOICENETTAMOUNT"]

used_orders = collections.Counter(); used_inv = collections.Counter(); sample = []
for f in sorted(glob.glob(OUT + "/*.xlsx")):
    code = [x for x in os.path.basename(f).split("_") if x.isdigit()][0]
    e = raw[code]; net = e[1] + e[3]
    wb = openpyxl.load_workbook(f)
    check(wb.sheetnames == ["Summary", "Customers by Region", "Online Orders", "Verification Notes"], f"{code}: sheets {wb.sheetnames}")
    ws = wb["Summary"]
    check(abs(ws["B5"].value - net) < 0.01, f"{code}: net {ws['B5'].value} != Ostendo {net}")
    check(ws["B6"].value == len(e[4]), f"{code}: customers {ws['B6'].value} != {len(e[4])}")
    rows = list(ws.iter_rows(min_row=9, max_row=200, values_only=True))
    chan = {r[0]: r for r in rows[:5] if r[0]}
    check(chan["Online"][1] + chan["Rep / Phone"][1] + chan["Corporate EDI"][1] == e[0], f"{code}: channel invoice count")
    check(abs(chan["Online"][2] + chan["Rep / Phone"][2] + chan["Corporate EDI"][2] - e[1]) < 0.01, f"{code}: channel revenue != Ostendo invoices")
    check(chan["Credit notes"][1] == e[2] and abs(chan["Credit notes"][2] - e[3]) < 0.01, f"{code}: credit notes")
    st = [i for i, r in enumerate(rows) if r[0] == "Region"][0]
    rr = []
    for r in rows[st + 1:]:
        if r[0] == "Total": tot = r; break
        rr.append(r)
    check(abs(sum(r[3] for r in rr) - net) < 0.01 and abs(tot[3] - net) < 0.01 and sum(r[2] for r in rr) == e[0] + e[2], f"{code}: region table does not add up")
    for ch in ("Online", "Rep / Phone", "Corporate EDI", "Credit notes"):
        check(abs(sum(r[3] for r in rr if r[1] == ch) - chan[ch][2]) < 0.01, f"{code}: region rows for {ch} != channel total")
    # customers sheet
    cs = [r for r in wb["Customers by Region"].iter_rows(min_row=2, values_only=True) if r[0] and r[0] != "Total"]
    check(len(cs) == len(e[4]) and abs(sum(r[3] for r in cs) - net) < 0.01, f"{code}: customers sheet total")
    check(abs(sum(r[8] for r in cs) - chan["Online"][2]) < 0.01 and abs(sum(r[9] for r in cs) - chan["Rep / Phone"][2]) < 0.01, f"{code}: customer online/rep columns")
    for r in cs:
        ch = [h for h in H if h["CUSTOMER"] == r[0] and h["SALESPERSON"].split("-")[0] == code and FROM <= pd(h["INVOICEDATE"]) <= TO]
        check(abs(sum(h["INVOICENETTAMOUNT"] for h in ch) - r[3]) < 0.01, f"{code}: customer {r[0]} revenue")
    # online sheet: every line proven against its source order
    on = [r for r in wb["Online Orders"].iter_rows(min_row=2, values_only=True) if r[1] and r[0] != "Total"]
    check(len(on) == chan["Online"][1] and abs(sum(r[5] for r in on) - chan["Online"][2]) < 0.01, f"{code}: online sheet != Online channel")
    for r in on:
        h = RAW.get(r[1]); used_inv[r[1]] += 1; sample.append(r[1])
        check(h and abs(h["INVOICENETTAMOUNT"] - r[5]) < 0.005 and h["SALESPERSON"].split("-")[0] == code and h["INVOICEORCREDIT"] == "Invoice", f"{code}: online line {r[1]} not an Ostendo invoice of this rep")
        if r[7].startswith("Same store, same amount, entered twice"):
            twin = [x for x in H if x["INVOICENUMBER"] != r[1] and x["DELIVERYCUSTOMER"] == h["DELIVERYCUSTOMER"] and x["INVOICEDATE"] == h["INVOICEDATE"] and abs(x["INVOICENETTAMOUNT"] - h["INVOICENETTAMOUNT"]) < 0.005]
            check(bool(twin), f"{code}: duplicate claim not true for {r[1]}"); continue
        used_orders[r[6]] += 1
        ev = EV.get(r[6])
        if not ev: fails.append(f"{code}: online order {r[6]} not found"); continue
        days = (pd(h["INVOICEDATE"]) - ev[0]).days
        if r[7].startswith("Same amount"):
            check(-1 <= days <= 14 and abs(r[5] - ev[1]) <= max(0.05, 0.002 * ev[1]), f"{code}: same-amount proof fails {r[1]} vs {r[6]} ({days}d, {r[5]} vs {ev[1]})")
        else:
            check(0 <= days <= 10 and 0.5 <= r[5] / ev[1] <= 1.6, f"{code}: re-key proof fails {r[1]} vs {r[6]} ({days}d, {r[5]} vs {ev[1]})")

check(all(v == 1 for v in used_orders.values()), "online order used twice: " + str([k for k, v in used_orders.items() if v > 1][:5]))
check(all(v == 1 for v in used_inv.values()), "invoice counted online twice")

random.seed(16)
pick = random.sample(sorted(set(sample)), min(25, len(set(sample)))) + random.sample(sorted(k for k, h in RAW.items() if FROM <= pd(h["INVOICEDATE"]) <= TO), 25)
ctx = ssl.create_default_context(); ctx.check_hostname = False; ctx.verify_mode = ssl.CERT_NONE
q = "SELECT INVOICENUMBER,INVOICEDATE,SALESPERSON,DELIVERYCUSTOMER,CUSTOMER,INVOICENETTAMOUNT FROM SALESINVOICEHEADER WHERE INVOICENUMBER IN (" + ",".join("'" + p + "'" for p in pick) + ")"
req = urllib.request.Request(f"{os.environ['OSTENDO_BASE_URL']}/sqlquery?apikey={os.environ['OSTENDO_API_KEY']}&format=json", data=q.encode(), headers={"Content-Type": "text/plain", "User-Agent": "curl/8.4.0"})
lv = {x["INVOICENUMBER"]: x for x in json.loads(urllib.request.urlopen(req, context=ctx).read())}
ok = 0
for p in pick:
    x, h = lv.get(p), RAW[p]
    if x and all(x[k] == h[k] for k in ("INVOICEDATE", "SALESPERSON", "DELIVERYCUSTOMER", "CUSTOMER")) and abs(x["INVOICENETTAMOUNT"] - h["INVOICENETTAMOUNT"]) < 0.005: ok += 1
    else: fails.append(f"live Ostendo differs for {p}")
print(f"live Ostendo spot check: {ok}/{len(pick)} identical")
print("FAILURES:", len(fails))
for x in fails[:30]: print("  ", x)
